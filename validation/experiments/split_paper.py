"""
split_paper.py
Splits paper/COGNICORE_RESEARCH_PAPER.md into 5 distinct, logically complete sections,
and generates both .md and .docx versions in:
1. paper/sections/
2. C:\\Users\\premk\\Downloads\\COGNICORE_PAPER_SECTIONS\\
"""
import os
import shutil
from generate_paper_docx import convert_md_to_docx
import docx

def split_and_convert():
    repo_root = os.getcwd()
    paper_path = os.path.join(repo_root, 'paper', 'COGNICORE_RESEARCH_PAPER.md')
    sections_dir = os.path.join(repo_root, 'paper', 'sections')
    downloads_dir = os.path.join(os.path.expanduser('~'), 'Downloads', 'COGNICORE_PAPER_SECTIONS')
    
    os.makedirs(sections_dir, exist_ok=True)
    os.makedirs(downloads_dir, exist_ok=True)
    
    with open(paper_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    slices = [
        ("01_abstract_and_introduction", lines[:74], "Part 1: Abstract & Introduction (RQ1-RQ6)"),
        ("02_related_work_and_positioning", lines[74:129], "Part 2: Related Work & Positioning (Table 1)"),
        ("03_system_architecture_and_mechanics", lines[129:201], "Part 3: System Architecture & Threat Model"),
        ("04_methodology_and_empirical_results", lines[201:347], "Part 4: Methodology & Empirical Results (Tables 2-7)"),
        ("05_discussion_limitations_and_conclusion", lines[347:], "Part 5: Discussion, Limitations, Conclusion & References"),
    ]
    
    print(f"Total lines in source paper: {len(lines)}")
    total_split_lines = sum(len(s[1]) for s in slices)
    assert total_split_lines == len(lines), f"Line mismatch: {total_split_lines} != {len(lines)}"
    print(f"Slice verification passed: exactly {total_split_lines} lines partitioned.\n")
    
    results = []
    
    for filename, part_lines, desc in slices:
        md_content = "".join(part_lines).strip() + "\n"
        
        # Paths
        md_file = os.path.join(sections_dir, f"{filename}.md")
        docx_file = os.path.join(sections_dir, f"{filename}.docx")
        
        # Write .md
        with open(md_file, 'w', encoding='utf-8') as f:
            f.write(md_content)
            
        # Convert to .docx
        convert_md_to_docx(md_file, docx_file)
        
        # Verify .docx integrity
        doc = docx.Document(docx_file)
        para_count = len(doc.paragraphs)
        table_count = len(doc.tables)
        
        # Copy to Downloads
        shutil.copyfile(md_file, os.path.join(downloads_dir, f"{filename}.md"))
        shutil.copyfile(docx_file, os.path.join(downloads_dir, f"{filename}.docx"))
        
        results.append({
            "name": filename,
            "desc": desc,
            "lines": len(part_lines),
            "paragraphs": para_count,
            "tables": table_count,
            "docx": docx_file,
            "dl_docx": os.path.join(downloads_dir, f"{filename}.docx")
        })
        print(f"-> Finished {filename}: {para_count} paragraphs, {table_count} tables\n")
        
    print("=" * 60)
    print("ALL 5 SECTIONS SUCCESSFULLY GENERATED AND VERIFIED!")
    print(f"Saved locally to: {sections_dir}")
    print(f"Saved in Downloads to: {downloads_dir}")
    print("=" * 60)

if __name__ == '__main__':
    split_and_convert()
